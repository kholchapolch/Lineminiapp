"""Read the Sony sampling workbook into a reviewable JSON dataset; never write DB.
Requires openpyxl. Explicitly records omitted rows and optional missing data.
"""
import argparse
import hashlib
import json
from datetime import datetime, date
from pathlib import Path
from urllib.parse import urlparse
import openpyxl

ALLOWED_URL_HOSTS = (
    'www.sony.co.th',
    'web.sony-asia.com',
    'experience.sony-asia.com',
    '8qym66.s.gy',
    'www.playstation.com',
)
FALLBACK_CONTENT_IMAGE_URL = 'https://www.sony.co.th/microsite/cs-portal/sony-logo.png'
CTA_ACTION_KEYS = {
    1: 'firmware',
    2: 'user-manual',
    3: 'recommended-apps',
    4: 'pro-tips',
    5: 'compatible-gear',
    6: 'view-my-badge',
    7: 'workshop',
    8: 'contest',
}

parser = argparse.ArgumentParser()
parser.add_argument('workbook')
parser.add_argument('--out', required=True)
args = parser.parse_args()
source = Path(args.workbook)
w = openpyxl.load_workbook(source, data_only=True)
products, contents, issues = [], [], []
article_stats = dict(content_id_key=0, synthetic_external_key=0)
carousel_stats = dict(carousel_id_key=0, synthetic_external_key=0)

def clean(v):
    return str(v).strip() if v is not None else ''

def issue(sheet, row, field, reason):
    issues.append(dict(sheet=sheet, row=row, field=field, reason=reason))

def url(v):
    v = clean(v)
    u = urlparse(v)
    if u.scheme != 'https' or u.hostname not in ALLOWED_URL_HOSTS or u.username or u.password:
        return None
    return v

def https_any_url(v):
    """Register CTA may point outside the Sony content allowlist (e.g. MySony staging)."""
    v = clean(v)
    u = urlparse(v)
    if u.scheme != 'https' or not u.hostname or u.username or u.password:
        return None
    return v

def https_image_url(v):
    v = clean(v)
    u = urlparse(v)
    if u.scheme != 'https' or not u.hostname or u.username or u.password:
        return None
    return v

def publish_time(v):
    if isinstance(v, datetime):
        return f"{v.strftime('%Y-%m-%d %H:%M:%S')}.{int(v.microsecond / 1000):03d}"
    if isinstance(v, date):
        return f"{v.isoformat()} 00:00:00.000"
    text = clean(v)
    if not text:
        return None
    for fmt in ('%Y-%m-%d %H:%M:%S.%f', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try:
            parsed = datetime.strptime(text, fmt)
            return f"{parsed.strftime('%Y-%m-%d %H:%M:%S')}.{int(parsed.microsecond / 1000):03d}"
        except ValueError:
            continue
    return None

def base(key, kind, locale='th'):
    return dict(external_key=key, locale=locale, content_type=kind, target_type='global', target_key='global', action_key=None, sort_order=0, published_at=None)

s = w['5. Product Image Data']
headers = [clean(v) for v in next(s.values)]
# Current workbook: Product_ID, Model Name, Category, Image URL, ...
# Older sampling workbook used Model/Image/Category order.
def product_fields(row):
    cells = list(row) + [None] * max(0, 4 - len(row))
    if 'Category' in headers and 'Image URL' in headers:
        return map(clean, (cells[0], cells[1], cells[3], cells[2]))
    return map(clean, cells[:4])

seen_models = set()
for n, row in enumerate(s.iter_rows(min_row=2, values_only=True), 2):
    if not any(v is not None for v in row): continue
    identifier, model, image, category = product_fields(row)
    if not model:
        issue(s.title,n,'Model_Name','ROW_EXCLUDED_MISSING_MODEL'); continue
    if 'ï¿½' in model or '\ufffd' in model:
        issue(s.title,n,'Model_Name','SOURCE_ENCODING_SUSPECT_PRESERVED_VERBATIM')
    model_key = model.upper()
    if model_key in seen_models:
        issue(s.title,n,'Model_Name','ROW_EXCLUDED_DUPLICATE_MODEL_KEY'); continue
    seen_models.add(model_key)
    image_url = https_image_url(image)
    if image_url and urlparse(image_url).path.rstrip('/') in ('','/th','/en'):
        image_url = None
    if not image_url: issue(s.title,n,'Model_Image_URL','IMAGE_UNAVAILABLE_STORED_NULL')
    category_code = category or None
    if category_code and (category_code.startswith('http://') or category_code.startswith('https://')):
        issue(s.title,n,'Category','CATEGORY_LOOKS_LIKE_URL_STORED_NULL'); category_code = None
    products.append(dict(external_key=identifier or None, model_name=model, model_key=model_key, category_code=category_code, image_url=image_url, sort_order=len(products)))

products_by_key = {p['model_key']: p for p in products}
mapping_stats = dict(
    added=0,
    skipped_existing=0,
    missing_kataban=0,
    missing_model=0,
    articles_cloned=0,
    carousels_cloned=0,
    content_skipped_duplicate=0,
    content_skipped_key_too_long=0,
    modelNames=[],
)
# MKT Kataban -> mapped Model Name keys (SKU variants that share image/content).
kataban_to_aliases = {}

# Expand catalog SKUs from Mobile/Game mapping: Model Name clones image/category from MKT Kataban.
mapping_sheet = '10. Mobile& Game Mapping'
if mapping_sheet not in w.sheetnames:
    issue('workbook', 0, mapping_sheet, 'SHEET_MISSING_MAPPING_EXPANSION_SKIPPED')
else:
    s = w[mapping_sheet]
    for n, row in enumerate(s.iter_rows(min_row=2, values_only=True), 2):
        if not any(v is not None for v in row):
            continue
        cells = list(row) + [None] * 4
        model = clean(cells[1])
        kataban = clean(cells[2])
        if not model:
            mapping_stats['missing_model'] += 1
            issue(s.title, n, 'Model_Name', 'ROW_EXCLUDED_MISSING_MODEL')
            continue
        if not kataban:
            mapping_stats['missing_kataban'] += 1
            issue(s.title, n, 'MKT_Kataban', 'ROW_EXCLUDED_MISSING_KATABAN')
            continue
        model_key = model.upper()
        kataban_key = kataban.upper()
        parent = products_by_key.get(kataban_key)
        if not parent:
            mapping_stats['missing_kataban'] += 1
            issue(s.title, n, 'MKT_Kataban', 'ROW_EXCLUDED_KATABAN_NOT_IN_PRODUCT_IMAGE')
            continue
        # Column B Model_Name successfully tied to a Product Image kataban.
        if model not in mapping_stats['modelNames']:
            mapping_stats['modelNames'].append(model)
        aliases = kataban_to_aliases.setdefault(kataban_key, [])
        if model_key not in aliases and model_key != kataban_key:
            aliases.append(model_key)
        if model_key in seen_models:
            mapping_stats['skipped_existing'] += 1
            issue(s.title, n, 'Model_Name', 'ROW_SKIPPED_MODEL_ALREADY_IN_CATALOG')
            continue
        seen_models.add(model_key)
        cloned = dict(
            external_key=None,
            model_name=model,
            model_key=model_key,
            category_code=parent['category_code'],
            image_url=parent['image_url'],
            sort_order=len(products),
        )
        products.append(cloned)
        products_by_key[model_key] = cloned
        mapping_stats['added'] += 1
        issue(s.title, n, 'Model_Name', 'MODEL_CLONED_FROM_KATABAN')

def ensure_model(model, sheet, row):
    model_key = model.upper()
    if model_key in seen_models:
        return model_key
    seen_models.add(model_key)
    products.append(dict(
        external_key=None,
        model_name=model,
        model_key=model_key,
        category_code=None,
        image_url=None,
        sort_order=len(products),
    ))
    issue(sheet, row, 'Model_Name', 'MODEL_STUB_ADDED_FROM_CONTENT')
    return model_key

# Articles/carousel: stable IDs when present; otherwise deterministic gen: hash.
s = w['6. Product Article Data']
headers = [clean(v) for v in next(s.values)]
seen_article_keys = set()
for n, row in enumerate(s.iter_rows(min_row=2, values_only=True), 2):
    if not any(v is not None for v in row):
        continue
    data = dict(zip(headers, row))
    model = clean(data.get('Model_Name'))
    title = clean(data.get('Title'))
    description = clean(data.get('Description')) or title
    cta_url = url(data.get('CTA_URL'))
    image_url = https_image_url(data.get('Thumbnail_Image_URL'))
    published_at = publish_time(data.get('PublishTime'))
    content_id = clean(data.get('Content_ID'))
    raw_cta = data.get('CTA_Index')
    try:
        cta_index = int(raw_cta) if raw_cta is not None and clean(raw_cta) != '' else None
    except (TypeError, ValueError):
        cta_index = None
    action_key = CTA_ACTION_KEYS.get(cta_index)

    if not image_url:
        image_url = FALLBACK_CONTENT_IMAGE_URL
        issue(s.title, n, 'Thumbnail_Image_URL', 'IMAGE_FALLBACK_USED')
    if description and description == title and not clean(data.get('Description')):
        issue(s.title, n, 'Description', 'DESCRIPTION_FALLBACK_TO_TITLE')

    missing = [name for name, value in [
        ('Model_Name', model),
        ('Title', title),
        ('Description', description),
        ('CTA_URL', cta_url),
        ('Thumbnail_Image_URL', image_url),
        ('PublishTime', published_at),
        ('CTA_Index', action_key),
    ] if not value]
    if missing:
        issue(s.title, n, ','.join(missing), 'ROW_EXCLUDED_MISSING_REQUIRED_FIELDS')
        continue

    model_key = ensure_model(model, s.title, n)

    if content_id:
        external_key = f'{content_id}:{model_key}'
        key_source = 'content_id_key'
    else:
        digest = hashlib.sha1(f'th|{action_key}|{model_key}|{cta_url}|{title}'.encode('utf-8')).hexdigest()[:16]
        external_key = f'gen:{digest}:{model_key}'
        key_source = 'synthetic_external_key'
        issue(s.title, n, 'Content_ID', 'SYNTHETIC_EXTERNAL_KEY')

    if len(external_key) > 191:
        issue(s.title, n, 'external_key', 'ROW_EXCLUDED_EXTERNAL_KEY_TOO_LONG')
        continue
    locale_type_key = f'th:article:{external_key}'.lower()
    if locale_type_key in seen_article_keys:
        issue(s.title, n, 'external_key', 'ROW_EXCLUDED_DUPLICATE_EXTERNAL_KEY')
        continue
    seen_article_keys.add(locale_type_key)

    article = base(external_key, 'article')
    article.update(
        target_type='model',
        target_key=model_key,
        action_key=action_key,
        sort_order=cta_index - 1,
        published_at=published_at,
        payload=dict(title=title, description=description, imageUrl=image_url, url=cta_url),
    )
    contents.append(article)
    article_stats[key_source] += 1

s = w['7. Product Carousel Data ']
headers = [clean(v) for v in next(s.values)]
seen_carousel_keys = set()
carousel_sort_by_model = {}
for n, row in enumerate(s.iter_rows(min_row=2, values_only=True), 2):
    if not any(v is not None for v in row):
        continue
    data = dict(zip(headers, row))
    model = clean(data.get('Model_Name'))
    title = clean(data.get('Title'))
    description = clean(data.get('Description')) or title
    cta_url = url(data.get('CTA_URL'))
    image_url = https_image_url(data.get('Thumbnail_Image_URL'))
    published_at = publish_time(data.get('PublishTime'))
    carousel_id = clean(data.get('Carousel_ID'))

    if not image_url:
        image_url = FALLBACK_CONTENT_IMAGE_URL
        issue(s.title, n, 'Thumbnail_Image_URL', 'IMAGE_FALLBACK_USED')
    if description and description == title and not clean(data.get('Description')):
        issue(s.title, n, 'Description', 'DESCRIPTION_FALLBACK_TO_TITLE')

    missing = [name for name, value in [
        ('Model_Name', model),
        ('Title', title),
        ('Description', description),
        ('CTA_URL', cta_url),
        ('Thumbnail_Image_URL', image_url),
        ('PublishTime', published_at),
    ] if not value]
    if missing:
        issue(s.title, n, ','.join(missing), 'ROW_EXCLUDED_MISSING_REQUIRED_FIELDS')
        continue

    model_key = ensure_model(model, s.title, n)

    if carousel_id:
        external_key = f'{carousel_id}:{model_key}'
        key_source = 'carousel_id_key'
    else:
        digest = hashlib.sha1(f'th|carousel|{model_key}|{cta_url}|{title}'.encode('utf-8')).hexdigest()[:16]
        external_key = f'gen:{digest}:{model_key}'
        key_source = 'synthetic_external_key'
        issue(s.title, n, 'Carousel_ID', 'SYNTHETIC_EXTERNAL_KEY')

    if len(external_key) > 191:
        issue(s.title, n, 'external_key', 'ROW_EXCLUDED_EXTERNAL_KEY_TOO_LONG')
        continue
    locale_type_key = f'th:carousel:{external_key}'.lower()
    if locale_type_key in seen_carousel_keys:
        issue(s.title, n, 'external_key', 'ROW_EXCLUDED_DUPLICATE_EXTERNAL_KEY')
        continue
    seen_carousel_keys.add(locale_type_key)

    sort_order = carousel_sort_by_model.get(model_key, 0)
    carousel_sort_by_model[model_key] = sort_order + 1

    item = base(external_key, 'carousel')
    item.update(
        target_type='model',
        target_key=model_key,
        action_key=None,
        sort_order=sort_order,
        published_at=published_at,
        payload=dict(title=title, description=description, imageUrl=image_url, url=cta_url),
    )
    contents.append(item)
    carousel_stats[key_source] += 1

# Clone model-targeted articles/carousels from each MKT Kataban onto mapped SKU Model Names.
seen_content_keys = {
    f"{row['locale']}:{row['content_type']}:{row['external_key']}".lower()
    for row in contents
}

def cloned_external_key(source_key, source_target, alias_key):
    suffix = f':{source_target}'
    if source_key.upper().endswith(suffix):
        return source_key[: -len(source_target)] + alias_key
    return f'{source_key}:{alias_key}'

expanded_contents = []
for row in contents:
    if row['content_type'] not in ('article', 'carousel') or row['target_type'] != 'model':
        continue
    aliases = kataban_to_aliases.get(row['target_key'])
    if not aliases:
        continue
    for alias_key in aliases:
        external_key = cloned_external_key(row['external_key'], row['target_key'], alias_key)
        if len(external_key) > 191:
            mapping_stats['content_skipped_key_too_long'] += 1
            issue(mapping_sheet, 0, 'external_key', 'CONTENT_CLONE_SKIPPED_EXTERNAL_KEY_TOO_LONG')
            continue
        locale_type_key = f"{row['locale']}:{row['content_type']}:{external_key}".lower()
        if locale_type_key in seen_content_keys:
            mapping_stats['content_skipped_duplicate'] += 1
            continue
        seen_content_keys.add(locale_type_key)
        cloned = dict(row)
        cloned['external_key'] = external_key
        cloned['target_key'] = alias_key
        cloned['payload'] = dict(row['payload'])
        expanded_contents.append(cloned)
        if row['content_type'] == 'article':
            mapping_stats['articles_cloned'] += 1
        else:
            mapping_stats['carousels_cloned'] += 1
contents.extend(expanded_contents)

s=w['8. Register Product Page Data']
for n,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
    if not any(v is not None for v in row): continue
    cells=list(row[:6])+[None]*(6-len(row[:6]))
    title,lead,image,heading,body,register = map(clean,cells)
    image_url=url(image)
    register_url=https_any_url(register) if register else None
    if not all((title,lead,heading,body,image_url)):
        issue(s.title,n,'page','ROW_EXCLUDED_INCOMPLETE_PAGE');continue
    blocks=[dict(type='image',url=image_url,alt=title),dict(type='heading',text=heading),dict(type='check_list',items=[line.strip().lstrip('-').strip() for line in body.splitlines() if line.strip()])]
    if register_url:
        blocks.append(dict(type='link_button',label=title,url=register_url))
    else:
        issue(s.title,n,'Register_Now_URL','REGISTER_URL_MISSING_LINK_BUTTON_OMITTED')
    c=base('register-product','page')
    c['payload']=dict(title=title,lead=lead,blocks=blocks)
    contents.append(c)

# CTA data lives in form controls; cell-only reading must not turn blanks into false.
# It is deliberately not imported until the CTA action/checkbox resolver review.
s=w['9. CTA Type (Omelet Mapping)']
labels={int(row[0]):(clean(row[1]),clean(row[2])) for row in s.iter_rows(min_row=3,values_only=True) if isinstance(row[0],(float,int))}
for n in range(3,11): issue(s.title,n,'checkbox/category/action','CTA_DEFERRED_TO_MAPPING_REVIEW')
linked_sheet='10. Linked URL Data'
if linked_sheet not in w.sheetnames:
    issue('workbook',0,linked_sheet,'SHEET_MISSING_FOOTER_LINKS_SKIPPED')
else:
    s=w[linked_sheet]
    for n,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
        if not any(v is not None for v in row):continue
        for col,key,index in [(0,'service-center',9),(1,'repair-status',10)]:
            link=url(row[col])
            if not link: issue(s.title,n,key,'ROW_EXCLUDED_INVALID_URL');continue
            for locale,label in zip(('en','th'),labels[index]):
                c=base(key,'footer_link',locale);c['sort_order']=col;c['payload']=dict(label=label,url=link);contents.append(c)
        issue(s.title,n,'FAQ_URL','FAQ_EXCLUDED_MISSING_URL' if not url(row[2]) else 'FAQ_LABEL_REQUIRES_REVIEW')

out=Path(args.out)
out.parent.mkdir(parents=True,exist_ok=True)
dataset=dict(products=products,contents=contents)
out.write_text(json.dumps(dataset,ensure_ascii=False,indent=2)+'\n')
report=dict(
    source=source.name,
    sourceSha256=hashlib.sha256(source.read_bytes()).hexdigest(),
    sampleData=True,
    policy='Import validated subset; Content_ID/Carousel_ID use ID:MODEL keys when present; missing IDs use deterministic gen: hash; article/carousel images accept any HTTPS host; missing thumbnails use sony-logo fallback; no invented PublishTime; missing product images are null; Mobile/Game mapping clones Product Image, article, and carousel rows for each Model Name under its MKT Kataban.',
    products=len(products),
    contents=len(contents),
    articles=article_stats,
    carousels=carousel_stats,
    mobileGameMapping=mapping_stats,
    issues=issues,
)
out.with_suffix('.report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
model_names_path = out.with_name('mobile-game-mapping-model-names.json')
model_names_path.write_text(
    json.dumps(
        dict(
            source='10. Mobile& Game Mapping Column B Model_Name',
            workbook=source.name,
            count=len(mapping_stats['modelNames']),
            modelNames=mapping_stats['modelNames'],
        ),
        ensure_ascii=False,
        indent=2,
    )
    + '\n'
)
print(json.dumps({k:v for k,v in report.items() if k!='issues'},ensure_ascii=False))
print('Issues:',len(issues))
print('Mobile/Game Column B model names:', model_names_path)
